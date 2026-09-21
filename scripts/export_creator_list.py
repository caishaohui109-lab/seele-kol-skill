#!/usr/bin/env python3
"""Export validated SEELE creator prospects to CSV and styled XLSX."""

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

REQUIRED_FIELDS = [
    "Creator ID",
    "达人名称",
    "平台",
    "主页链接",
    "国家/语言",
    "粉丝数",
    "近期相关内容样本量",
    "中位播放",
    "内容标签",
    "代表内容链接",
    "公开商务邮箱",
    "建议触达渠道",
    "达人类型",
    "UGC交付能力",
    "游戏创作动机",
    "受众重合度",
    "产品说服力",
    "单位传播效率",
    "商务可执行性",
    "推荐拍摄方式",
    "推荐内容切角",
    "具体测试场景",
    "实验假设",
    "测试价值",
    "与现有达人同质性",
    "推荐等级",
    "推荐理由",
    "风险或待确认项",
    "人工复核重点",
    "个性化私信草稿",
    "报价USD",
    "预估CPM",
    "数据核验日期",
    "人工确认状态",
    "触达状态",
]

GRADE_COLORS = {
    "A": "C6EFCE",
    "B": "FFEB9C",
    "C": "E7E6E6",
    "Pass": "FFC7CE",
}


def validate_rows(rows):
    for index, row in enumerate(rows, start=2):
        missing = [field for field in REQUIRED_FIELDS if field not in row]
        if missing:
            raise ValueError(f"row {index} missing fields: {', '.join(missing)}")


def export_rows(rows, output_dir, date_stamp):
    validate_rows(rows)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = f"SEELE_达人触达名单_{date_stamp}"
    csv_path = output_dir / f"{stem}.csv"
    xlsx_path = output_dir / f"{stem}.xlsx"

    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=REQUIRED_FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "达人触达名单"
    sheet.append(REQUIRED_FIELDS)
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    grade_col = REQUIRED_FIELDS.index("推荐等级") + 1
    for row in rows:
        sheet.append([row[field] for field in REQUIRED_FIELDS])
        color = GRADE_COLORS.get(row["推荐等级"])
        if color:
            sheet.cell(sheet.max_row, grade_col).fill = PatternFill(
                "solid", fgColor=color
            )
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
        sheet.column_dimensions[letter].width = max(width, 12)
    workbook.save(xlsx_path)
    return csv_path, xlsx_path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input_csv")
    parser.add_argument("output_dir")
    parser.add_argument("date_stamp")
    args = parser.parse_args()
    with Path(args.input_csv).open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    print(export_rows(rows, args.output_dir, args.date_stamp))


if __name__ == "__main__":
    main()
